import io
import pytest
from fastapi import HTTPException
from unittest.mock import MagicMock
import numpy as np

from ingestion.extractor import (
    extract_from_text,
    extract_from_docx,
    extract_from_pdf,
    extract_document,
)
from engine.document_db import DocumentDB
from clients.ollama_client import OllamaClient
from api.services.document_service import insert_document, ingest_file


# ── Extractor Tests ──────────────────────────────────────────────────────────

def test_extract_from_text_plain():
    content = b"This is a test markdown document.\n\n# Heading 1\nContent."
    sections = extract_from_text(content)
    assert len(sections) == 1
    assert "Heading 1" in sections[0].text
    assert sections[0].page_number is None


def test_extract_from_text_empty():
    content = b"   \n  \t  "
    sections = extract_from_text(content)
    assert len(sections) == 0


def test_extract_dispatcher_unsupported():
    with pytest.raises(ValueError, match="Unsupported file format"):
        extract_document("test.exe", b"binary content")


def test_extract_dispatcher_txt():
    sections = extract_document("notes.txt", b"Hello world!")
    assert len(sections) == 1
    assert sections[0].text == "Hello world!"


def test_extract_docx_mock():
    import docx
    bio = io.BytesIO()
    doc = docx.Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    doc.save(bio)
    bio.seek(0)

    sections = extract_from_docx(bio.getvalue())
    assert len(sections) == 1
    assert "First paragraph." in sections[0].text
    assert "Second paragraph." in sections[0].text


def test_extract_pdf_mock():
    import pypdf
    bio = io.BytesIO()
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(bio)
    bio.seek(0)

    # Empty page should extract safely to 0 non-empty text sections
    sections = extract_from_pdf(bio.getvalue())
    assert isinstance(sections, list)


# ── Ingest Service & Error Handling Tests ────────────────────────────────────

def test_insert_document_empty_error():
    doc_db = DocumentDB()
    mock_ollama = MagicMock(spec=OllamaClient)
    mock_ollama.is_available.return_value = True

    with pytest.raises(HTTPException) as exc:
        insert_document(doc_db, mock_ollama, "EmptyDoc", "   ")
    assert exc.value.status_code == 422


def test_insert_document_ollama_unavailable():
    doc_db = DocumentDB()
    mock_ollama = MagicMock(spec=OllamaClient)
    mock_ollama.is_available.return_value = False

    with pytest.raises(HTTPException) as exc:
        insert_document(doc_db, mock_ollama, "Doc", "Valid text")
    assert exc.value.status_code == 503


def test_insert_document_rollback_on_failure():
    doc_db = DocumentDB()
    mock_ollama = MagicMock(spec=OllamaClient)
    mock_ollama.is_available.return_value = True
    # First chunk succeeds, second chunk fails
    mock_ollama.embed.side_effect = [np.zeros(16, dtype=np.float32), None]

    long_text = " ".join([f"word{i}" for i in range(400)])
    with pytest.raises(HTTPException) as exc:
        insert_document(doc_db, mock_ollama, "Doc", long_text)
    assert exc.value.status_code == 502
    # Ensure rolled back completely (no residual items)
    assert doc_db.size == 0
