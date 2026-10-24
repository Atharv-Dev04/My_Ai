"""
Text extraction module for multi-format document ingestion.

Supported formats:
    - PDF (.pdf) via pypdf (extracts page text, preserves 1-indexed page numbers where available)
    - DOCX (.docx) via python-docx (paragraphs and tables)
    - TXT (.txt) UTF-8 plain text with fallback decoders
    - Markdown (.md, .markdown) UTF-8 plain text

All extractors produce a list of ExtractedSection dataclasses preserving source metadata.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from typing import BinaryIO


@dataclass
class ExtractedSection:
    """Represents text extracted from a document with optional page attribution."""
    text: str
    page_number: int | None = None


def extract_from_pdf(content: bytes | BinaryIO) -> list[ExtractedSection]:
    """Extract text page-by-page from a PDF file."""
    import pypdf

    stream = io.BytesIO(content) if isinstance(content, bytes) else content
    try:
        reader = pypdf.PdfReader(stream)
    except Exception as e:
        raise ValueError(f"Malformed or unreadable PDF: {e}") from e

    sections: list[ExtractedSection] = []
    for idx, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
            cleaned = page_text.strip()
            if cleaned:
                sections.append(ExtractedSection(text=cleaned, page_number=idx))
        except Exception:
            # Skip unreadable single page, continue with rest
            continue

    return sections


def extract_from_docx(content: bytes | BinaryIO) -> list[ExtractedSection]:
    """Extract text from a Word DOCX file (paragraphs and table cells)."""
    import docx

    stream = io.BytesIO(content) if isinstance(content, bytes) else content
    try:
        doc = docx.Document(stream)
    except Exception as e:
        raise ValueError(f"Malformed or unreadable DOCX: {e}") from e

    paragraphs: list[str] = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            paragraphs.append(txt)

    for table in doc.tables:
        for row in table.rows:
            row_txt = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_txt:
                paragraphs.append(" | ".join(row_txt))

    full_text = "\n\n".join(paragraphs).strip()
    if not full_text:
        return []

    # DOCX does not expose paginated physical page boundaries
    return [ExtractedSection(text=full_text, page_number=None)]


def extract_from_text(content: bytes) -> list[ExtractedSection]:
    """Extract text from plain text or markdown files with encoding fallback."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            decoded = content.decode(encoding).strip()
            if decoded:
                return [ExtractedSection(text=decoded, page_number=None)]
            return []
        except UnicodeDecodeError:
            continue

    raise ValueError("File content could not be decoded as text.")


def extract_document(filename: str, content: bytes) -> list[ExtractedSection]:
    """
    Dispatcher based on file extension.

    Raises:
        ValueError for unsupported extensions or corrupt files.
    """
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_from_pdf(content)
    elif lower.endswith(".docx"):
        return extract_from_docx(content)
    elif lower.endswith((".txt", ".md", ".markdown")):
        return extract_from_text(content)
    else:
        raise ValueError(
            f"Unsupported file format '{filename}'. "
            f"Supported formats: .pdf, .docx, .txt, .md, .markdown"
        )
